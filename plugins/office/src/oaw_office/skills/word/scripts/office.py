import argparse
from itertools import islice

from docx import Document
from docx.text.paragraph import Paragraph

from office_runtime import clipped, emit_main, office_input, save_new, spec, window


def main():
    parser = argparse.ArgumentParser(description="Read, create or edit DOCX paragraphs and tables")
    commands = parser.add_subparsers(dest="command", required=True)
    read = commands.add_parser("inspect")
    read.add_argument("--input", required=True)
    window(read)
    for name in ("create", "edit"):
        command = commands.add_parser(name)
        command.add_argument("--spec", required=True)
        command.add_argument("--output", required=True)
        if name == "edit":
            command.add_argument("--input", required=True)
    args = parser.parse_args()
    document = Document(office_input(args.input, ".docx")) if hasattr(args, "input") else Document()
    if args.command == "inspect":
        if args.offset < 0 or not 1 <= args.limit <= args.window_max:
            raise ValueError("Use offset >= 0 and limit between 1 and 200")
        total = len(document.paragraphs) + len(document.tables)
        result, truncated = [], False
        paragraph_index, table_index = -1, -1
        for index, block in enumerate(document.iter_inner_content()):
            if isinstance(block, Paragraph):
                paragraph_index += 1
            else:
                table_index += 1
            if index < args.offset:
                continue
            if index >= args.offset + args.limit:
                break
            if isinstance(block, Paragraph):
                text, cut = clipped(block.text)
                result.append({"block": index, "index": paragraph_index, "kind": "paragraph", "text": text,
                    "style": block.style.name, "text_truncated": cut})
                truncated |= cut
            else:
                rows = []
                for row in islice(block.rows, 100):
                    cells = [clipped(cell.text) for cell in row.cells[:50]]
                    rows.append([text for text, _ in cells])
                    truncated |= any(cut for _, cut in cells) or len(row.cells) > 50
                cut = len(block.rows) > 100
                truncated |= cut
                result.append({"block": index, "index": table_index, "kind": "table", "rows": rows,
                    "total_rows": len(block.rows), "rows_truncated": cut})
        next_offset = args.offset + len(result)
        return {"blocks": result, "total_blocks": total, "next_offset": next_offset if next_offset < total else None,
            "truncated": truncated or next_offset < total,
            "paragraph_count": len(document.paragraphs), "table_count": len(document.tables),
            "warnings": ["Main-body text and tables only; headers, footers, revisions, drawings and page layout were not inspected."]}
    payload = spec(args.spec)
    if args.command == "create":
        for block in payload["blocks"]:
            if block["type"] == "paragraph":
                document.add_paragraph(block["text"], style=block.get("style"))
            elif block["type"] == "table":
                rows = block["rows"]
                width = max(len(row) for row in rows)
                table = document.add_table(rows=0, cols=width)
                table.style = "Table Grid"
                for row in rows:
                    cells = table.add_row().cells
                    for column, value in enumerate(row):
                        cells[column].text = str(value)
            else:
                raise ValueError("Block type must be paragraph or table")
    else:
        for change in payload.get("paragraphs", []):
            index = change["index"]
            if not isinstance(index, int) or not 0 <= index < len(document.paragraphs):
                raise ValueError("Paragraph index is out of range")
            document.paragraphs[index].text = change["text"]
        for change in payload.get("cells", []):
            indices = (change["table"], change["row"], change["column"])
            if any(not isinstance(index, int) or index < 0 for index in indices):
                raise ValueError("Table, row and column indices must be nonnegative integers")
            table = document.tables[indices[0]]
            if indices[1] >= len(table.rows) or indices[2] >= len(table.columns):
                raise ValueError("Table cell is out of range")
            table.cell(indices[1], indices[2]).text = change["text"]
    result = save_new(document, args.output, ".docx", Document)
    result["warnings"] = ["Text replacement resets inline formatting in the changed paragraph or cell; inspect formatting before delivery."] if args.command == "edit" else []
    return result


if __name__ == "__main__":
    emit_main(main)
