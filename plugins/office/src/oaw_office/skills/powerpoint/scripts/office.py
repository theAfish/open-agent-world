import argparse
from itertools import islice

from pptx import Presentation
from pptx.util import Inches, Pt

from office_runtime import clipped, emit_main, office_input, save_new, spec, window


def slide_at(presentation, number):
    if not isinstance(number, int) or not 1 <= number <= len(presentation.slides):
        raise ValueError("Slide number is out of range (slides start at 1)")
    return presentation.slides[number - 1]


def main():
    parser = argparse.ArgumentParser(description="Inspect PPTX slides and notes, create or edit presentations")
    commands = parser.add_subparsers(dest="command", required=True)
    read = commands.add_parser("inspect")
    read.add_argument("--input", required=True)
    window(read, default=5, maximum=20)
    for name in ("create", "edit"):
        command = commands.add_parser(name)
        command.add_argument("--spec", required=True)
        command.add_argument("--output", required=True)
        if name == "edit":
            command.add_argument("--input", required=True)
    args = parser.parse_args()
    presentation = Presentation(office_input(args.input, ".pptx")) if hasattr(args, "input") else Presentation()
    warnings = ["Embedded images, charts, grouped shapes, animations and visual layout were not inspected."]
    if args.command == "inspect":
        if args.offset < 0 or not 1 <= args.limit <= args.window_max:
            raise ValueError("Use offset >= 0 and limit between 1 and 20")
        slides, cut = [], False
        for index, slide in enumerate(islice(presentation.slides, args.offset, args.offset + args.limit), args.offset + 1):
            shapes = []
            cut |= len(slide.shapes) > 200
            for shape in islice(slide.shapes, 200):
                item = {"shape_id": shape.shape_id, "name": shape.name, "type": str(shape.shape_type)}
                if shape.has_text_frame:
                    item["text"], text_cut = clipped(shape.text)
                    item["text_truncated"] = text_cut
                    cut |= text_cut
                if shape.has_table:
                    item["table"] = []
                    for row in islice(shape.table.rows, 50):
                        cells = [clipped(cell.text) for cell in islice(row.cells, 30)]
                        item["table"].append([text for text, _ in cells])
                        cut |= any(truncated for _, truncated in cells) or len(row.cells) > 30
                    cut |= len(shape.table.rows) > 50
                shapes.append(item)
            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text
            notes, notes_cut = clipped(notes, 4000)
            cut |= notes_cut
            slides.append({"slide": index, "shapes": shapes, "notes": notes, "notes_truncated": notes_cut})
        next_offset = args.offset + len(slides)
        return {"slides": slides, "total_slides": len(presentation.slides),
            "next_offset": next_offset if next_offset < len(presentation.slides) else None,
            "truncated": cut or next_offset < len(presentation.slides), "warnings": warnings}
    payload = spec(args.spec)
    if args.command == "create":
        presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
        for definition in payload["slides"]:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = definition["title"]
            slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)
            frame = slide.placeholders[1].text_frame
            for index, text in enumerate(definition.get("bullets", [])):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text, paragraph.font.size = text, Pt(22)
            if "notes" in definition:
                slide.notes_slide.notes_text_frame.text = definition["notes"]
    else:
        for change in payload.get("text", []):
            slide = slide_at(presentation, change["slide"])
            shape = next((shape for shape in slide.shapes if shape.shape_id == change["shape_id"]), None)
            if shape is None or not shape.has_text_frame:
                raise ValueError("Shape must identify an inspected top-level text frame")
            shape.text = change["text"]
        for change in payload.get("notes", []):
            slide_at(presentation, change["slide"]).notes_slide.notes_text_frame.text = change["text"]
        warnings.append("Text replacement resets inline formatting in changed shapes.")
    result = save_new(presentation, args.output, ".pptx", Presentation)
    result["warnings"] = warnings
    return result


if __name__ == "__main__":
    emit_main(main)
