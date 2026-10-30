import argparse
import re

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from office_runtime import clipped, emit_main, office_input, save_new, spec


def main():
    parser = argparse.ArgumentParser(description="Read formulas and cached values, create or edit XLSX workbooks")
    commands = parser.add_subparsers(dest="command", required=True)
    read = commands.add_parser("inspect")
    read.add_argument("--input", required=True)
    read.add_argument("--sheet")
    read.add_argument("--start-row", type=int, default=1)
    read.add_argument("--rows", type=int, default=100)
    read.add_argument("--start-column", type=int, default=1)
    read.add_argument("--columns", type=int, default=20)
    for name in ("create", "edit"):
        command = commands.add_parser(name)
        command.add_argument("--spec", required=True)
        command.add_argument("--output", required=True)
        if name == "edit":
            command.add_argument("--input", required=True)
    args = parser.parse_args()
    warnings = ["Formulas are not recalculated. Cached results may be missing or stale.",
                "Charts, drawings and print layout have not been visually checked."]
    if args.command == "inspect":
        if not (1 <= args.rows <= 200 and 1 <= args.columns <= 50
                and 1 <= args.start_row <= 1048576 and 1 <= args.start_column <= 16384):
            raise ValueError("Use 1..200 rows, 1..50 columns and valid positive Excel coordinates")
        path = office_input(args.input, ".xlsx")
        workbook = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[args.sheet] if args.sheet else workbook.active
            values = cached[sheet.title]
            end_row = min(sheet.max_row, args.start_row + args.rows - 1)
            end_col = min(sheet.max_column, args.start_column + args.columns - 1)
            rows, cut, missing = [], False, 0
            raw_rows = sheet.iter_rows(min_row=args.start_row, max_row=end_row, min_col=args.start_column, max_col=end_col)
            value_rows = values.iter_rows(min_row=args.start_row, max_row=end_row, min_col=args.start_column, max_col=end_col)
            for row_index, (raw, stored) in enumerate(zip(raw_rows, value_rows), args.start_row):
                cells = []
                for col_index, (cell, value) in enumerate(zip(raw, stored), args.start_column):
                    item = {"address": f"{get_column_letter(col_index)}{row_index}", "value": value.value}
                    if cell.data_type == "f":
                        item.update(formula=cell.value, cache_missing=value.value is None)
                        missing += value.value is None
                    for key in ("value", "formula"):
                        if isinstance(item.get(key), str):
                            item[key], text_cut = clipped(item[key])
                            if text_cut:
                                item[key + "_truncated"] = True
                                cut = True
                    cells.append(item)
                rows.append(cells)
            return {"sheet": sheet.title, "sheets": [{"name": s.title, "state": s.sheet_state,
                "rows": s.max_row, "columns": s.max_column} for s in workbook.worksheets],
                "rows": rows, "missing_formula_caches": missing,
                "next_row": end_row + 1 if end_row < sheet.max_row else None,
                "next_column": end_col + 1 if end_col < sheet.max_column else None,
                "truncated": cut or end_row < sheet.max_row or end_col < sheet.max_column,
                "warnings": warnings}
        finally:
            workbook.close()
            cached.close()
    payload = spec(args.spec)
    if args.command == "create":
        if not payload["sheets"]:
            raise ValueError("Provide at least one worksheet")
        workbook = Workbook()
        workbook.remove(workbook.active)
        for definition in payload["sheets"]:
            if definition["name"] in workbook.sheetnames:
                raise ValueError("Worksheet names must be unique")
            sheet = workbook.create_sheet(definition["name"])
            for row in definition["rows"]:
                sheet.append(row)
            sheet.freeze_panes = "A2"
    else:
        workbook = load_workbook(office_input(args.input, ".xlsx"), data_only=False, rich_text=True)
        for change in payload["cells"]:
            if not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}", change["cell"]):
                raise ValueError("Use a single A1-style cell address")
            cell = workbook[change["sheet"]][change["cell"]]
            cell.value = change["value"]
            if "number_format" in change:
                cell.number_format = change["number_format"]
    try:
        result = save_new(workbook, args.output, ".xlsx", load_workbook)
        result["warnings"] = warnings
        return result
    finally:
        workbook.close()


if __name__ == "__main__":
    emit_main(main)
